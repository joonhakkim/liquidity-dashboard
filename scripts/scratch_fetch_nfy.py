import json, sys, time
import pandas as pd
import openpyxl
sys.path.insert(0, '.')
from build_op_band import (
    HEADER_CODE_ROW, HEADER_ITEM_ROW, HEADER_BASEDATE_ROW, DATA_START_ROW,
    QUARTER_ITEM, find_workbook, detect_blocks,
)

with open('../scratch_bottom_all.json', encoding='utf-8') as f:
    rows = json.load(f)
names = [r[1] for r in rows]

sm = pd.read_csv('../data/sector_map.csv', dtype={'code': str})
name_to_code = dict(zip(sm['name'], sm['code']))
name_to_sector = dict(zip(sm['name'], sm['sector']))
name_to_code['티씨머트리얼즈'] = '125020'
name_to_sector['티씨머트리얼즈'] = '전기장비'

wb_path = find_workbook()
print('workbook:', wb_path, flush=True)
t0 = time.time()
wb = openpyxl.load_workbook(wb_path, read_only=True, data_only=True)
print('loaded in', round(time.time() - t0, 1), 's', flush=True)

wanted_codes = {('A' + name_to_code[n]) for n in names if n in name_to_code}
print('wanted', len(wanted_codes), flush=True)

nfy_data = {}
for sn in wb.sheetnames:
    if len(nfy_data) >= len(wanted_codes):
        break
    ws = wb[sn]
    max_col = ws.max_column
    t1 = time.time()

    def read_row(r):
        return next(ws.iter_rows(min_row=r, max_row=r, max_col=max_col, values_only=True))

    row_codes = read_row(HEADER_CODE_ROW)
    row_items = read_row(HEADER_ITEM_ROW)
    row_basedate = read_row(HEADER_BASEDATE_ROW)
    print(f'{sn}: header read in', round(time.time() - t1, 1), 's', flush=True)
    blocks = detect_blocks(row_codes)
    relevant = [b for b in blocks if b[0] in wanted_codes and b[0] not in nfy_data]
    if not relevant:
        continue
    print(f'{sn}: {len(relevant)} matches out of {len(blocks)} blocks', flush=True)

    for code, start, end in relevant:
        nfy1_idx = nfy2_idx = None
        for k in range(start, end):
            item, base = row_items[k], row_basedate[k]
            if item == QUARTER_ITEM:
                if base == 'NFY1':
                    nfy1_idx = k
                elif base == 'NFY2':
                    nfy2_idx = k
        if nfy1_idx is None or nfy2_idx is None:
            nfy_data[code] = (None, None)
            continue
        lo, hi = sorted([nfy1_idx, nfy2_idx])
        fy1_last = fy2_last = None
        for row in ws.iter_rows(min_row=DATA_START_ROW, max_row=ws.max_row,
                                 min_col=lo + 1, max_col=hi + 1, values_only=True):
            v1 = row[nfy1_idx - lo]
            v2 = row[nfy2_idx - lo]
            if v1 is not None:
                fy1_last = v1
            if v2 is not None:
                fy2_last = v2
        nfy_data[code] = (fy1_last, fy2_last)
    print(f'{sn}: done, elapsed', round(time.time() - t1, 1), 's, resolved so far', len(nfy_data), flush=True)

with open('../scratch_nfy.json', 'w', encoding='utf-8') as f:
    json.dump({'name_to_code': name_to_code, 'name_to_sector': name_to_sector, 'nfy_data': nfy_data}, f, ensure_ascii=False)
print('done, resolved', len(nfy_data), flush=True)
