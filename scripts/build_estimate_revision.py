"""
이익추정치 상향 트래커 - "기업 밴드 찾기 2" 신형 파일(NFY1/NFY2 컬럼, 매일 갱신)에서
당해년도(FY1)/차년도(FY2) 영업이익 컨센서스가 최근 상향되고 있는 종목을 찾는다.

OP밴드 트래커와는 목적이 달라서 별도 페이지로 분리했다(2026-09-07 사용자 요청) - OP밴드는
"6월 스위칭 룰로 지금 어느 연도 추정치를 쓸지 골라서 배수가 얼마인지" 보는 거고, 이건
"NFY1/NFY2 각각을 스위칭 없이 그대로 연속 시계열로 두고 그 값 자체가 시간에 따라 상향/
하향되고 있는지"를 본다. 그래서 구형(분기 합산) 파일은 지원 안 하고 신형 파일 전용이다 -
구형은애초에 이 정도로 자주 안 바뀌어서 "상향 추이" 자체가 의미가 약함.

산출 로직:
1) build_op_band.py와 동일한 방식으로 회사별 컬럼블록(시가총액/TTM/NFY1/NFY2)을 탐지한다.
   NFY1/NFY2가 없는 종목(구형 파일에만 있던 분기형 종목)은 건너뛴다.
2) 날짜별 NFY1/NFY2 값을 스위칭 없이 그대로 시계열로 모은다(값이 없는 날은 결측으로 둠 -
   회계연도가 안 채워진 기간을 이전 값으로 채우거나 하지 않는다. 지어내지 않는다).
3) 최신값 대비 "N 거래일 전" 값과 비교해서 상향률(%)을 1일/1주(5거래일)/1개월(21거래일)
   세 가지로 계산한다. 두 값 다 있어야 계산하고, 과거값이 0이면(혹은 결측이면) None.
4) 결과를 종목별 JSON으로 페이지에 통째로 임베드한다(OP밴드처럼 종목당 파일로 안 쪼갬 -
   시계열 전체가 아니라 요약 숫자 몇 개만 필요해서 용량이 가볍다).
"""
import json
import os
from datetime import datetime

import openpyxl
import pandas as pd

from build_op_band import (
    MKTCAP_ITEM, TTM_ITEM, QUARTER_ITEM, HEADER_CODE_ROW, HEADER_NAME_ROW,
    HEADER_ITEM_ROW, HEADER_BASEDATE_ROW, DATA_START_ROW,
    find_workbook, detect_blocks, load_naver_sector_map, load_sector_map,
)

DATA_DIR = os.path.join(os.path.dirname(__file__), "..", "data")
DOCS_DIR = os.path.join(os.path.dirname(__file__), "..", "docs")
OUT_PATH = os.path.join(DOCS_DIR, "estimate_revision.html")

LOOKBACKS = [("1d", 1, "1일"), ("1w", 5, "1주"), ("1m", 21, "1개월")]


def process_sheet(ws):
    max_col = ws.max_column
    max_row = ws.max_row

    def read_row(r):
        return next(ws.iter_rows(min_row=r, max_row=r, max_col=max_col, values_only=True))

    row_codes = read_row(HEADER_CODE_ROW)
    row_names = read_row(HEADER_NAME_ROW)
    row_items = read_row(HEADER_ITEM_ROW)
    row_basedate = read_row(HEADER_BASEDATE_ROW)

    blocks = detect_blocks(row_codes)

    dates = []
    data_rows = []
    for row in ws.iter_rows(min_row=DATA_START_ROW, max_row=max_row, max_col=max_col, values_only=True):
        d = row[0]
        if d is None:
            continue
        if isinstance(d, datetime):
            dates.append(d)
            data_rows.append(row)

    results = {}
    for code, start, end in blocks:
        mktcap_idx = ttm_idx = nfy1_idx = nfy2_idx = None
        for k in range(start, end):
            item = row_items[k]
            base = row_basedate[k]
            if item == MKTCAP_ITEM:
                mktcap_idx = k
            elif item == TTM_ITEM:
                ttm_idx = k
            elif item == QUARTER_ITEM:
                if base == "NFY1":
                    nfy1_idx = k
                elif base == "NFY2":
                    nfy2_idx = k

        # 신형 파일 전용 - NFY1/NFY2 둘 다 없는(구형 분기형) 종목은 건너뛴다.
        if mktcap_idx is None or (nfy1_idx is None and nfy2_idx is None):
            continue

        name = row_names[start] if start < len(row_names) else code
        series_dates, series_mktcap, series_fy1, series_fy2, series_ttm = [], [], [], [], []
        for row_vals, d in zip(data_rows, dates):
            mktcap = row_vals[mktcap_idx]
            if mktcap is None:
                continue
            fy1 = row_vals[nfy1_idx] if nfy1_idx is not None else None
            fy2 = row_vals[nfy2_idx] if nfy2_idx is not None else None
            ttm = row_vals[ttm_idx] if ttm_idx is not None else None
            if fy1 is None and fy2 is None and ttm is None:
                continue
            series_dates.append(d.strftime("%Y-%m-%d"))
            series_mktcap.append(mktcap)
            series_fy1.append(fy1)
            series_fy2.append(fy2)
            series_ttm.append(ttm)

        if not series_dates:
            continue
        results[code] = {"name": name, "dates": series_dates, "mktcap": series_mktcap,
                          "fy1": series_fy1, "fy2": series_fy2, "ttm": series_ttm}
    return results


def pct_change(series, back):
    """series의 마지막 값 대비 back개 데이터포인트 전 값의 변화율(%). 둘 중 하나라도 없거나
    과거값이 0이면 None(지어내지 않음)."""
    if len(series) <= back:
        return None
    latest = series[-1]
    past = series[-1 - back]
    if latest is None or past is None or past == 0:
        return None
    return round((latest / past - 1) * 100, 2)


def build_summary_row(code, data, sector_map, naver_sector_map):
    dates = data["dates"]
    fy1 = data["fy1"]
    fy2 = data["fy2"]
    ttm = data["ttm"]
    mktcap = data["mktcap"]

    # FY2(차년도, 27년) 컨센서스 자체가 없는 종목이 많다("27년은 없는거니까") - 그런 종목은
    # TTM(최근 12개월 실적)을 대신 비교 대상으로 쓴다(2026-09-07 사용자 요청). FY1도 같은
    # 논리로 폴백하되, 실제로 FY1이 비는 경우는 드물다. 폴백 여부는 fy1_source/fy2_source에
    # "FY"/"TTM"으로 남겨서 화면에 표시("TTM 비교"라고 티 나게) - 어느 걸로 비교했는지 숨기지
    # 않는다.
    latest_fy1 = fy1[-1]
    fy1_source = "FY"
    fy1_series = fy1
    if latest_fy1 is None:
        latest_fy1 = ttm[-1]
        fy1_source = "TTM"
        fy1_series = ttm

    latest_fy2 = fy2[-1]
    fy2_source = "FY"
    fy2_series = fy2
    if latest_fy2 is None:
        latest_fy2 = ttm[-1]
        fy2_source = "TTM"
        fy2_series = ttm

    if latest_fy1 is None and latest_fy2 is None:
        return None

    bare_code = code.lstrip("A")
    sector = naver_sector_map.get(bare_code) or sector_map.get(data["name"])

    row = {
        "code": code, "name": data["name"], "sector": sector,
        "latest_date": dates[-1],
        "mktcap": round(mktcap[-1], 0) if mktcap[-1] is not None else None,
        "fy1": round(latest_fy1, 0) if latest_fy1 is not None else None,
        "fy2": round(latest_fy2, 0) if latest_fy2 is not None else None,
        "fy1_source": fy1_source, "fy2_source": fy2_source,
    }
    for key, back, _ in LOOKBACKS:
        row[f"fy1_{key}"] = pct_change(fy1_series, back)
        row[f"fy2_{key}"] = pct_change(fy2_series, back)

    # FY1(당해) 대비 FY2(차년도) 성장률 - "추정치가 시간이 지나며 상향/하향됐는지"(위 컬럼들)와는
    # 다른 지표로, "지금 시점에 내년이 올해보다 얼마나 더 좋아질 걸로 보는지"(2026-09-07 요청).
    # 둘 중 하나라도 TTM 폴백값이면 순수 FY1/FY2 비교가 아니라는 걸 growth_mixed_ttm으로 표시.
    if latest_fy1 not in (None, 0) and latest_fy2 is not None:
        row["fy2_vs_fy1_growth"] = round((latest_fy2 / latest_fy1 - 1) * 100, 2)
    else:
        row["fy2_vs_fy1_growth"] = None
    row["growth_mixed_ttm"] = fy1_source == "TTM" or fy2_source == "TTM"
    return row


TEMPLATE = """<!doctype html>
<html lang="ko">
<head>
<meta charset="utf-8">
<title>이익추정치 상향 트래커</title>
<style>
  body {{ font-family: -apple-system, "Malgun Gothic", sans-serif; background:#0f1115; color:#e6e6e6; margin:0; padding:24px; }}
  a.back {{ color:#4dabf7; font-size:13px; text-decoration:none; }}
  h1 {{ font-size:20px; margin:8px 0 4px 0; }}
  .updated {{ color:#9aa0a6; font-size:13px; margin-bottom:20px; }}
  .filters {{ display:flex; flex-wrap:wrap; gap:14px; align-items:center; margin-bottom:16px;
    background:#1a1d24; border:1px solid #23262e; border-radius:10px; padding:14px 16px; }}
  .filters label {{ font-size:13px; color:#9aa0a6; display:flex; align-items:center; gap:6px; }}
  .filters input, .filters select {{ background:#0f1115; border:1px solid #23262e; color:#e6e6e6;
    border-radius:6px; padding:5px 8px; font-size:13px; }}
  .filters input[type="text"] {{ width:110px; }}
  .filters input[type="number"] {{ width:80px; }}
  table {{ width:100%; border-collapse:collapse; font-size:13px; }}
  th, td {{ padding:8px 10px; text-align:right; border-bottom:1px solid #23262e; white-space:nowrap; }}
  th:nth-child(1), td:nth-child(1), th:nth-child(2), td:nth-child(2), th:nth-child(3), td:nth-child(3) {{ text-align:left; }}
  th {{ color:#9aa0a6; font-weight:600; position:sticky; top:0; background:#0f1115; cursor:pointer; user-select:none; }}
  th:hover {{ color:#e6e6e6; }}
  tbody tr:hover {{ background:#1a1d24; cursor:pointer; }}
  .up {{ color:#ff6b6b; }}
  .down {{ color:#4dabf7; }}
  .ttm-tag {{ color:#6b7280; font-size:10px; border:1px solid #23262e; border-radius:4px; padding:1px 4px; }}
  .table-wrap {{ overflow-x:auto; border:1px solid #23262e; border-radius:10px; }}
  .hint {{ color:#6b7280; font-size:12px; margin:10px 0 0 0; }}
</style>
</head>
<body>
  <a class="back" href="index.html">&larr; 홈</a>
  <h1>이익추정치 상향 트래커</h1>
  <div class="updated">최종 갱신: {updated_at} &middot; 원본 파일 기준일 {src_mtime} &middot; {n_stocks}종목 &middot;
    당해({fy1_year}년)/차년({fy2_year}년) 영업이익 추정치 - 매일 갱신되는 컨센서스 값 그대로, 회계연도 스위칭 없음</div>

  <div class="filters">
    <label>검색 <input type="text" id="fSearch" placeholder="종목명/코드"></label>
    <label>섹터 <select id="fSector"><option value="">전체</option></select></label>
    <label>기준기간
      <select id="fPeriod">
        <option value="1d">1일</option>
        <option value="1w" selected>1주</option>
        <option value="1m">1개월</option>
      </select>
    </label>
    <label>정렬
      <select id="fSort">
        <option value="fy1_desc">{fy1_year}년 상향률 높은순</option>
        <option value="fy2_desc">{fy2_year}년 상향률 높은순</option>
        <option value="fy1_asc">{fy1_year}년 하향률 큰순</option>
        <option value="fy2_asc">{fy2_year}년 하향률 큰순</option>
        <option value="growth_desc">{fy1_year}년→{fy2_year}년 성장률 높은순</option>
        <option value="growth_asc">{fy1_year}년→{fy2_year}년 성장률 낮은순</option>
      </select>
    </label>
    <label>시가총액 최소(억원) <input type="number" id="fMktcapMin" step="100"></label>
  </div>

  <div class="table-wrap">
  <table>
    <thead><tr>
      <th>종목명</th><th>코드</th><th>섹터</th><th>시가총액</th>
      <th>{fy1_year}년 추정치</th><th data-sort="fy1">{fy1_year}년 상향률</th>
      <th>{fy2_year}년 추정치</th><th data-sort="fy2">{fy2_year}년 상향률</th>
      <th>{fy1_year}년&rarr;{fy2_year}년 성장률</th>
      <th>기준일</th>
    </tr></thead>
    <tbody id="tbody"></tbody>
  </table>
  </div>
  <div class="hint">상향률 = (최신값 / N거래일 전 값 - 1) x 100. 데이터가 그만큼 안 쌓인 종목은 "-"로 표시됩니다.
    <span class="ttm-tag">TTM</span> 표시는 해당 연도 컨센서스 자체가 없어서 최근 12개월 실적(TTM)으로 대신 비교했다는 뜻입니다.
    {fy1_year}년&rarr;{fy2_year}년 성장률은 "지금 시점 컨센서스로 내년이 올해보다 얼마나 좋아질 것으로 보는지"이고, 위 상향률(추정치가 최근 얼마나 바뀌었는지)과는 다른 지표입니다.</div>

<script>
const ROWS = {rows_json};

function fmtMktcap(won) {{
  if (won === null || won === undefined) return '-';
  const eok = won / 1e8;
  const jo = Math.floor(eok / 10000);
  const rest = Math.round(eok - jo * 10000);
  if (jo > 0) return `${{jo}}조 ${{rest.toLocaleString()}}억`;
  return `${{Math.round(eok).toLocaleString()}}억`;
}}
function fmtOp(v, source) {{
  if (v === null || v === undefined) return '-';
  const txt = fmtMktcap(v * 1000); // 원본 단위: 천원 -> 억 표기 재사용 위해 x1000
  // FY 추정치 자체가 없어서 TTM(최근 12개월 실적)으로 대신 비교한 경우 표시(2026-09-07) -
  // 어떤 기준으로 비교했는지 안 가리고 그대로 보여준다.
  return source === 'TTM' ? `${{txt}} <span class="ttm-tag">TTM</span>` : txt;
}}
function fmtPct(v) {{
  if (v === null || v === undefined) return '-';
  const cls = v > 0 ? 'up' : (v < 0 ? 'down' : '');
  const sign = v > 0 ? '+' : '';
  return `<span class="${{cls}}">${{sign}}${{v.toFixed(1)}}%</span>`;
}}

function applyFilters() {{
  const search = document.getElementById('fSearch').value.trim().toLowerCase();
  const sector = document.getElementById('fSector').value;
  const period = document.getElementById('fPeriod').value;
  const sort = document.getElementById('fSort').value;
  const mktcapMin = parseFloat(document.getElementById('fMktcapMin').value);

  let rows = ROWS.filter(r => {{
    if (search && !r.name.toLowerCase().includes(search) && !r.code.toLowerCase().includes(search)) return false;
    if (sector && r.sector !== sector) return false;
    if (!isNaN(mktcapMin) && (r.mktcap === null || r.mktcap / 1e8 < mktcapMin)) return false;
    return true;
  }});

  const fy1key = 'fy1_' + period, fy2key = 'fy2_' + period;
  const [sortField, dir] = sort.split('_');
  const sortKey = sortField === 'fy1' ? fy1key : sortField === 'fy2' ? fy2key : 'fy2_vs_fy1_growth';
  rows.sort((a, b) => {{
    const av = a[sortKey], bv = b[sortKey];
    if (av === null && bv === null) return 0;
    if (av === null) return 1;
    if (bv === null) return -1;
    return (av - bv) * (dir === 'desc' ? -1 : 1);
  }});

  const tbody = document.getElementById('tbody');
  tbody.innerHTML = rows.map(r => `
    <tr>
      <td>${{r.name}}</td>
      <td>${{r.code}}</td>
      <td>${{r.sector ?? '-'}}</td>
      <td>${{fmtMktcap(r.mktcap)}}</td>
      <td>${{fmtOp(r.fy1, r.fy1_source)}}</td>
      <td>${{fmtPct(r[fy1key])}}</td>
      <td>${{fmtOp(r.fy2, r.fy2_source)}}</td>
      <td>${{fmtPct(r[fy2key])}}</td>
      <td>${{fmtPct(r.fy2_vs_fy1_growth)}}${{r.growth_mixed_ttm && r.fy2_vs_fy1_growth !== null ? ' <span class="ttm-tag">TTM포함</span>' : ''}}</td>
      <td>${{r.latest_date}}</td>
    </tr>
  `).join('');
}}

const sectorSelect = document.getElementById('fSector');
[...new Set(ROWS.map(r => r.sector).filter(Boolean))].sort().forEach(s => {{
  const opt = document.createElement('option');
  opt.value = s; opt.textContent = s;
  sectorSelect.appendChild(opt);
}});

['fSearch', 'fSector', 'fPeriod', 'fSort', 'fMktcapMin'].forEach(id => {{
  document.getElementById(id).addEventListener('input', applyFilters);
  document.getElementById(id).addEventListener('change', applyFilters);
}});

applyFilters();
</script>
</body>
</html>
"""


def main():
    wb_path = find_workbook()
    if not wb_path:
        print("data/manual/ 에 '기업 밴드' 워크북이 없습니다.")
        return

    print(f"워크북 로드 중: {wb_path}")
    wb = openpyxl.load_workbook(wb_path, read_only=True, data_only=True)

    all_results = {}
    for sn in wb.sheetnames:
        print(f"처리 중: {sn} ...")
        results = process_sheet(wb[sn])
        for code, data in results.items():
            if code not in all_results:
                all_results[code] = data
        print(f"  {len(results)}개 종목(신형 NFY1/NFY2 보유) 처리")

    print(f"\n총 {len(all_results)}개 종목")

    naver_sector_map = load_naver_sector_map()
    sector_map = load_sector_map()
    print(f"네이버 업종 매핑 {len(naver_sector_map)}종목, 큐레이션 섹터 매핑 {len(sector_map)}종목 로드됨")

    rows = []
    for code, data in all_results.items():
        row = build_summary_row(code, data, sector_map, naver_sector_map)
        if row is not None:
            rows.append(row)

    # 커버리지가 끊긴(최신값이 오래된) 종목은 뺀다 - 예전 값끼리 비교하면 "몇 년 전 대비
    # +5000%" 같은 의미 없는 상향률이 나온다(2026-09-07 확인). 전체에서 가장 최근 기준일을
    # 가진 종목들만 남긴다 - 매일 갱신되는 파일이라 살아있는 커버리지는 다 같은 날짜여야 함.
    latest_overall = max((r["latest_date"] for r in rows), default=None)
    before = len(rows)
    rows = [r for r in rows if r["latest_date"] == latest_overall]
    print(f"결과 {before}종목 중 최신 기준일({latest_overall}) 종목만 유지: {len(rows)}종목 "
          f"(커버리지 끊긴 {before - len(rows)}종목 제외)")

    # NFY1은 "그 날짜가 속한 연도", NFY2는 "다음 연도"로 확인됨(2026-09-03) - 데이터의 실제
    # 최신 기준일에서 연도를 뽑아온다(빌드 실행 시각이 아니라 데이터 기준일 기준이어야 연말/
    # 연초 경계에서도 어긋나지 않음).
    fy1_year = int(latest_overall[:4]) if latest_overall else datetime.now().year
    fy2_year = fy1_year + 1

    src_mtime = datetime.fromtimestamp(os.path.getmtime(wb_path)).strftime("%Y-%m-%d %H:%M")
    html = TEMPLATE.format(
        updated_at=datetime.now().strftime("%Y-%m-%d %H:%M"),
        src_mtime=src_mtime,
        n_stocks=len(rows),
        fy1_year=fy1_year,
        fy2_year=fy2_year,
        rows_json=json.dumps(rows, ensure_ascii=False),
    )
    os.makedirs(DOCS_DIR, exist_ok=True)
    with open(OUT_PATH, "w", encoding="utf-8") as f:
        f.write(html)
    print(f"저장 완료: {OUT_PATH}")


if __name__ == "__main__":
    main()
