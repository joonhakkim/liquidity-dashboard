# -*- coding: utf-8 -*-
import json
import openpyxl

with open('../scratch_nfy.json', encoding='utf-8') as f:
    d = json.load(f)
name_to_code = d['name_to_code']
name_to_sector = d['name_to_sector']
nfy_data = d['nfy_data']

path = r'C:\매크로파일\바텀종목\바텀 종목 구성 2026-09-28.xlsx'
wb = openpyxl.load_workbook(path)
ws = wb[wb.sheetnames[0]]

missing = []
filled = 0
for r in range(3, 73):
    name = ws.cell(row=r, column=2).value
    if name is None:
        continue
    code = name_to_code.get(name)
    sector = name_to_sector.get(name)
    if sector is not None:
        ws.cell(row=r, column=6).value = sector
    if code is None:
        missing.append((r, name, 'code not found'))
        continue
    a_code = 'A' + code
    pair = nfy_data.get(a_code)
    if not pair or pair[0] is None or pair[1] is None or pair[0] == 0:
        missing.append((r, name, f'nfy missing: {pair}'))
        continue
    fy1, fy2 = pair
    growth = round((fy2 / fy1 - 1) * 100, 1)
    ws.cell(row=r, column=5).value = growth
    filled += 1

wb.save(path)
print('filled', filled)
print('missing:')
for m in missing:
    print(' ', m)
