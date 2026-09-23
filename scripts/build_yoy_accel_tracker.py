"""
YoY 가속화 트래커 - docs/yoy_accel_tracker.html (2026-09-23 사용자 요청)

data/manual/QoQ 계산.xlsx(FnGuide DataGuide 내보내기, "기업 밴드 찾기" 파일과 같은 회사블록
포맷 - 행8=코드/9=이름/10=Item Code/12=기준일자, 15행부터 일별 데이터)에서 종목별
시가총액(S102100, 일별)과 분기 영업이익(M121500.M=실적, E121500.M=추정, 각 분기 컬럼에
YYYYMM 라벨이 헤더12행에 붙어 있음)을 뽑아, 전년동기대비(YoY) 영업이익 증감률이 "가속"되고
있는 종목을 스크리닝한다.

컨센서스(E121500.M, 애널리스트 추정치)가 하나라도 있는 종목만 포함한다 - 사용자 요청
("컨센서스가 있는 종목들만 선택해서") 그대로. 추정치가 전혀 없으면 이 트래커가 보려는
"앞으로도 이어질 성장세인지"를 가늠할 근거 자체가 없다.

처음엔 시가총액과 함께 차트로 보여줬는데, 종목을 클릭했을 때 숫자 자체를 바로 보고 싶다는
요청(2026-09-23, "그냥 숫자자체를 보여줄 수 있게 그래프는 지워도 좋아")으로 차트를 빼고
분기별 시기/영업이익/YoY 표로 바꿨다.

부호가 섞이면(적자<->흑자 전환) 단순 비율이 정반대로 오해를 부르므로(OP밴드 v2/이익추정치
상향 트래커와 동일 원칙) 그 구간은 YoY%를 계산하지 않고 None으로 둔다.
"""
import glob
import json
import os
from datetime import datetime

import openpyxl
import pandas as pd

from build_op_band import (
    HEADER_CODE_ROW, HEADER_NAME_ROW, HEADER_ITEM_ROW, HEADER_BASEDATE_ROW, DATA_START_ROW,
    MANUAL_DIR, detect_blocks, load_naver_sector_map, load_sector_map,
)

DATA_DIR = os.path.join(os.path.dirname(__file__), "..", "data")
DOCS_DIR = os.path.join(os.path.dirname(__file__), "..", "docs")
SCREEN_DIR = os.path.join(DATA_DIR, "screening")
SUMMARY_PATH = os.path.join(SCREEN_DIR, "yoy_accel_summary.csv")
PAGE_OUT_PATH = os.path.join(DOCS_DIR, "yoy_accel_tracker.html")
DETAIL_OUT_DIR = os.path.join(DOCS_DIR, "yoy_accel_tracker_data")

MKTCAP_ITEM = "S102100"
OP_ACTUAL_ITEM = "M121500.M"
OP_EST_ITEM = "E121500.M"
MIN_QUARTERS_FOR_YOY = 5  # 최소 5분기(작년 동기 1개 비교 가능) 있어야 대상에 포함


def find_workbook():
    candidates = glob.glob(os.path.join(MANUAL_DIR, "*QoQ*.xls*"))
    candidates = [c for c in candidates if not os.path.basename(c).startswith("~$")]
    if not candidates:
        return None
    return max(candidates, key=os.path.getmtime)


def process_sheet(ws):
    max_col = ws.max_column
    max_row = ws.max_row

    def read_row(r):
        return next(ws.iter_rows(min_row=r, max_row=r, max_col=max_col, values_only=True))

    row_codes = read_row(HEADER_CODE_ROW)
    row_names = read_row(HEADER_NAME_ROW)
    row_items = read_row(HEADER_ITEM_ROW)
    row_basedate = read_row(HEADER_BASEDATE_ROW)

    blocks = detect_blocks(row_codes)

    dates, data_rows = [], []
    for row in ws.iter_rows(min_row=DATA_START_ROW, max_row=max_row, max_col=max_col, values_only=True):
        d = row[0]
        if d is None or not isinstance(d, datetime):
            continue
        dates.append(d)
        data_rows.append(row)

    results = {}
    for code, start, end in blocks:
        mktcap_idx = None
        quarter_cols = []  # (col_idx, period_yyyymm, is_estimate)
        for k in range(start, end):
            item = row_items[k]
            if item == MKTCAP_ITEM:
                mktcap_idx = k
            elif item in (OP_ACTUAL_ITEM, OP_EST_ITEM):
                period = row_basedate[k]
                if period is not None:
                    quarter_cols.append((k, int(period), item == OP_EST_ITEM))
        if mktcap_idx is None or not quarter_cols:
            continue

        name = row_names[start] if start < len(row_names) else code

        # 표(숫자) 위주 페이지라 시가총액은 요약표의 참고용 "최신값"만 있으면 된다(2026-09-23
        # 사용자 요청 - 차트를 표로 대체, 시계열 전체는 더 이상 안 씀).
        latest_mktcap = None
        for row_vals, d in zip(data_rows, dates):
            v = row_vals[mktcap_idx]
            if v is not None:
                latest_mktcap = v

        quarters = []
        for col_idx, period, is_est in sorted(quarter_cols, key=lambda x: x[1]):
            val = None
            for row_vals in data_rows:
                v = row_vals[col_idx]
                if v is not None:
                    val = v  # 분기 내내 동일값이 반복되므로 마지막(=가장 최근 갱신) 값을 쓴다
            if val is not None:
                quarters.append({"period": period, "is_estimate": is_est, "op_won": round(val * 1000, 0)})

        if len(quarters) < MIN_QUARTERS_FOR_YOY:
            continue
        has_consensus = any(q["is_estimate"] for q in quarters)
        if not has_consensus:
            continue

        results[code] = {
            "name": name, "latest_mktcap": latest_mktcap, "quarters": quarters,
        }
    return results


def compute_yoy(quarters):
    """quarters: [{period, is_estimate, op_won}] 기간순 정렬됨. 4분기(=1년) 전과 비교해
    YoY% 리스트를 같은 길이로 반환(앞 4개는 None). 부호가 섞이면(적자<->흑자) None."""
    yoy = [None] * len(quarters)
    for i in range(4, len(quarters)):
        cur = quarters[i]["op_won"]
        prev = quarters[i - 4]["op_won"]
        if cur is None or prev is None or prev == 0:
            continue
        if (cur > 0) != (prev > 0):
            continue
        yoy[i] = round((cur / prev - 1) * 100, 1)
    return yoy


def build_row_and_detail(code, data, sector_map, naver_sector_map):
    quarters = data["quarters"]
    yoy = compute_yoy(quarters)
    if not any(v is not None for v in yoy):
        return None, None

    # 최신 YoY / 직전 YoY -> 가속 여부(YoY 자체가 전분기보다 더 높아지고 있는지)
    idxs_with_yoy = [i for i, v in enumerate(yoy) if v is not None]
    latest_i = idxs_with_yoy[-1]
    latest_yoy = yoy[latest_i]
    latest_q = quarters[latest_i]
    prev_yoy = yoy[idxs_with_yoy[-2]] if len(idxs_with_yoy) >= 2 and idxs_with_yoy[-2] == latest_i - 1 else None
    accel = round(latest_yoy - prev_yoy, 1) if prev_yoy is not None else None

    # 세자리 YoY(100%+) 유지 분기 수 - 최신 분기부터 거꾸로 훑어서 100% 이상이 끊기지 않고
    # 몇 분기째 이어지는지(2026-09-23 사용자 요청 "세자리 YoY가 3분기 이상 유지되는 걸로
    # 필터"). 중간에 YoY가 없거나(부호 전환 등) 100% 밑으로 내려가면 그 자리에서 스트릭이
    # 끊긴다.
    triple_digit_streak = 0
    for i in range(latest_i, -1, -1):
        if yoy[i] is not None and yoy[i] >= 100:
            triple_digit_streak += 1
        else:
            break

    # 가속화(전분기 대비 YoY 상승) 유지 분기 수 - 최신 분기부터 거꾸로, 바로 앞 분기 대비
    # YoY가 계속 더 높아지고 있는 구간이 몇 분기째 이어지는지(2026-09-23 사용자 요청
    # "전분기 대비 가속화가 3분기 이상 유지되고 있는 애들"). 중간에 YoY가 비거나(부호전환 등)
    # 바로 직전 분기가 아니면(연속이 아니면) 비교 자체가 안 되므로 그 자리에서 끊긴다.
    accel_streak = 0
    i = latest_i
    while i - 1 >= 0 and yoy[i] is not None and yoy[i - 1] is not None and yoy[i] > yoy[i - 1]:
        accel_streak += 1
        i -= 1

    row = {
        "code": code, "name": data["name"],
        "sector": naver_sector_map.get(code.lstrip("A")) or sector_map.get(data["name"]),
        "latest_mktcap": data["latest_mktcap"],
        "latest_period": latest_q["period"], "latest_is_estimate": latest_q["is_estimate"],
        "latest_yoy": latest_yoy, "prev_yoy": prev_yoy, "accel": accel,
        "n_quarters": len(quarters), "triple_digit_streak": triple_digit_streak,
        "accel_streak": accel_streak,
    }

    # 클릭 시 보여줄 표 - 시기/영업이익/YoY만(2026-09-23 사용자 요청, 차트 대신 숫자 표).
    detail = {
        "code": code, "name": data["name"],
        "quarters": [
            {"period": q["period"], "op_100mil": round(q["op_won"] / 1e8, 1),
             "yoy": yoy[i], "is_estimate": q["is_estimate"]}
            for i, q in enumerate(quarters)
        ],
    }
    return row, detail


def main():
    wb_path = find_workbook()
    if not wb_path:
        print("data/manual/ 에 'QoQ 계산' 워크북이 없습니다.")
        return
    print(f"워크북 로드 중: {wb_path}")
    wb = openpyxl.load_workbook(wb_path, read_only=True, data_only=True)

    all_results = {}
    for sn in wb.sheetnames:
        print(f"처리 중: {sn} ...")
        res = process_sheet(wb[sn])
        for code, data in res.items():
            all_results.setdefault(code, data)
        print(f"  {len(res)}개 종목(컨센서스 보유, 5분기 이상)")

    naver_sector_map = load_naver_sector_map()
    sector_map = load_sector_map()

    os.makedirs(DETAIL_OUT_DIR, exist_ok=True)
    rows = []
    for code, data in all_results.items():
        row, detail = build_row_and_detail(code, data, sector_map, naver_sector_map)
        if not row:
            continue
        rows.append(row)
        with open(os.path.join(DETAIL_OUT_DIR, f"{code}.json"), "w", encoding="utf-8") as f:
            json.dump(detail, f, ensure_ascii=False)

    print(f"결과 {len(rows)}종목(YoY 계산 가능 + 컨센서스 보유)")

    os.makedirs(SCREEN_DIR, exist_ok=True)
    pd.DataFrame(rows).to_csv(SUMMARY_PATH, index=False, encoding="utf-8-sig")

    html = TEMPLATE.format(
        updated_at=datetime.now().strftime("%Y-%m-%d %H:%M"),
        n_stocks=len(rows),
        rows_json=json.dumps(rows, ensure_ascii=False),
    )
    os.makedirs(DOCS_DIR, exist_ok=True)
    with open(PAGE_OUT_PATH, "w", encoding="utf-8") as f:
        f.write(html)
    print(f"저장 완료: {SUMMARY_PATH}")
    print(f"저장 완료: {PAGE_OUT_PATH}")


TEMPLATE = """<!doctype html>
<html lang="ko">
<head>
<meta charset="utf-8">
<title>YoY 가속화 트래커</title>
<style>
  body {{ font-family: -apple-system, "Malgun Gothic", sans-serif; background:#0f1115; color:#e6e6e6; margin:0; padding:24px; }}
  a.back {{ color:#4dabf7; font-size:13px; text-decoration:none; margin-right:12px; }}
  h1 {{ font-size:20px; margin:8px 0 4px 0; }}
  .updated {{ color:#9aa0a6; font-size:13px; margin-bottom:16px; }}
  .exp {{ background:#1a1d24; border-radius:10px; padding:14px 18px; font-size:12px; color:#9aa0a6; line-height:1.8; max-width:1000px; margin-bottom:18px; }}
  .exp b {{ color:#ffa94d; }}
  .filters {{ display:flex; gap:12px; flex-wrap:wrap; align-items:center; margin-bottom:14px; font-size:13px; color:#9aa0a6; }}
  .filters input, .filters select {{ background:#1a1d24; border:1px solid #2a2e37; color:#e6e6e6; border-radius:6px; padding:6px 10px; font-size:13px; }}
  .filters input[type="number"] {{ width:80px; }}
  .table-wrap {{ overflow-x:auto; }}
  table {{ border-collapse:collapse; width:100%; font-size:13px; }}
  th, td {{ padding:8px 12px; text-align:right; border-bottom:1px solid #23262e; white-space:nowrap; }}
  th:first-child, td:first-child, th:nth-child(3), td:nth-child(3) {{ text-align:left; }}
  th {{ color:#9aa0a6; font-weight:normal; font-size:12px; position:sticky; top:0; background:#0f1115; }}
  .up {{ color:#63e6be; }}
  .down {{ color:#ff8787; }}
  .est-badge {{ display:inline-block; background:#a9c8ec33; color:#4dabf7; border:1px solid #4dabf7; border-radius:4px; font-size:10px; padding:1px 4px; margin-left:5px; }}
  .sub {{ color:#6b7280; font-size:11px; }}
  .count {{ color:#63e6be; font-size:12px; margin-bottom:8px; }}
  tbody tr {{ cursor:pointer; }}
  tbody tr:hover {{ background:#1a1d24; }}
  .overlay {{ display:none; position:fixed; inset:0; background:rgba(0,0,0,0.75); z-index:200; overflow:auto; padding:40px 20px; }}
  .overlay.open {{ display:block; }}
  .modal {{ background:#12151b; border:1px solid #23262e; border-radius:14px; max-width:520px; margin:0 auto; padding:22px 26px; }}
  .modal h2 {{ font-size:17px; margin:0 0 12px 0; }}
  .close-btn {{ float:right; background:none; border:none; color:#9aa0a6; font-size:22px; cursor:pointer; line-height:1; }}
  .detail-table {{ width:100%; }}
  .detail-table th {{ position:static; }}
</style>
</head>
<body>
  <a class="back" href="index.html">&larr; 홈</a>
  <h1>YoY 가속화 트래커</h1>
  <div class="updated">최종 갱신: {updated_at} &middot; {n_stocks}종목(컨센서스 보유)</div>

  <div class="exp">
    <b>무엇을 보는 페이지인가</b><br>
    분기별 영업이익(실적+애널리스트 컨센서스 추정치)의 <b>전년동기대비(YoY) 증감률</b>과,
    그 YoY 증감률 자체가 <b>전분기보다 더 가속되고 있는지</b>(가속도 = 이번 분기 YoY − 전분기 YoY)를
    같이 봅니다. 종목명을 클릭하면 분기별 시기/영업이익/YoY% 표를 볼 수 있습니다.
    적자/흑자가 뒤바뀌는 구간은 YoY%가 왜곡되므로 계산하지 않습니다(표에서 제외).
    FnGuide 컨센서스(추정치)가 하나도 없는 종목은 애초에 포함하지 않습니다.
  </div>

  <div class="filters">
    <label>검색 <input type="text" id="fSearch" placeholder="종목명/코드"></label>
    <label>섹터 <select id="fSector"><option value="">전체</option></select></label>
    <label>최신 YoY% 최소 <input type="number" id="fYoyMin" step="10"></label>
    <label><input type="checkbox" id="fAccelOnly"> 가속 중인 종목만(YoY가 전분기보다 상승)</label>
    <label><input type="checkbox" id="fTripleOnly"> 세자리 YoY(100%+) 3분기 이상 유지</label>
    <label><input type="checkbox" id="fAccelStreakOnly"> 가속화 3분기 이상 유지(전분기 대비 YoY 계속 상승)</label>
    <label>시총 최소(억) <input type="number" id="fMktcapMin" step="100"></label>
    <label>정렬 <select id="fSort">
      <option value="accel_desc">가속도 큰순</option>
      <option value="yoy_desc">최신 YoY% 큰순</option>
      <option value="streak_desc">세자리 유지 분기수 큰순</option>
      <option value="accel_streak_desc">가속화 유지 분기수 큰순</option>
      <option value="mktcap_desc">시가총액 큰순</option>
    </select></label>
  </div>
  <div class="count" id="count"></div>

  <div class="table-wrap">
  <table>
    <thead><tr>
      <th>종목명</th><th>코드</th><th>섹터</th><th>시가총액</th><th>최신 분기</th><th>최신 YoY%</th><th>전분기 YoY%</th><th>가속도(%p)</th><th>세자리 유지</th><th>가속화 유지</th>
    </tr></thead>
    <tbody id="tbody"></tbody>
  </table>
  </div>

  <div class="overlay" id="overlay">
    <div class="modal">
      <button class="close-btn" id="closeBtn">&times;</button>
      <h2 id="detailName"></h2>
      <table class="detail-table">
        <thead><tr><th style="text-align:left;">시기</th><th>영업이익(억원)</th><th>YoY%</th></tr></thead>
        <tbody id="detailBody"></tbody>
      </table>
    </div>
  </div>

<script>
const ROWS = {rows_json};

const sectors = [...new Set(ROWS.map(r => r.sector).filter(Boolean))].sort();
const selSector = document.getElementById('fSector');
sectors.forEach(s => {{ const o = document.createElement('option'); o.value = s; o.textContent = s; selSector.appendChild(o); }});

function fmtMktcap(v) {{
  if (v == null) return '-';
  const eok = v / 1e8;
  return eok >= 10000 ? (eok / 10000).toFixed(2) + '조' : Math.round(eok).toLocaleString() + '억';
}}
function fmtPeriod(p) {{
  const y = Math.floor(p / 100), m = p % 100;
  const q = {{3:'Q1',6:'Q2',9:'Q3',12:'Q4'}}[m] || m;
  return y + ' ' + q;
}}
function pctSpan(v) {{
  if (v == null) return '<span class="sub">-</span>';
  const cls = v > 0 ? 'up' : v < 0 ? 'down' : '';
  return `<span class="${{cls}}">${{v > 0 ? '+' : ''}}${{v.toFixed(1)}}%</span>`;
}}

function applyFilters() {{
  const q = document.getElementById('fSearch').value.trim().toLowerCase();
  const sec = selSector.value;
  const yoyMin = parseFloat(document.getElementById('fYoyMin').value);
  const mktcapMin = parseFloat(document.getElementById('fMktcapMin').value);
  const accelOnly = document.getElementById('fAccelOnly').checked;
  const tripleOnly = document.getElementById('fTripleOnly').checked;
  const accelStreakOnly = document.getElementById('fAccelStreakOnly').checked;
  const sort = document.getElementById('fSort').value;

  let rows = ROWS.filter(r => {{
    if (q && !(r.name.toLowerCase().includes(q) || r.code.toLowerCase().includes(q))) return false;
    if (sec && r.sector !== sec) return false;
    if (!isNaN(yoyMin) && (r.latest_yoy == null || r.latest_yoy < yoyMin)) return false;
    if (!isNaN(mktcapMin) && (r.latest_mktcap == null || r.latest_mktcap / 1e8 < mktcapMin)) return false;
    if (accelOnly && (r.accel == null || r.accel <= 0)) return false;
    if (tripleOnly && r.triple_digit_streak < 3) return false;
    if (accelStreakOnly && r.accel_streak < 3) return false;
    return true;
  }});

  if (sort === 'accel_desc') rows.sort((a, b) => (b.accel ?? -9e9) - (a.accel ?? -9e9));
  else if (sort === 'yoy_desc') rows.sort((a, b) => (b.latest_yoy ?? -9e9) - (a.latest_yoy ?? -9e9));
  else if (sort === 'streak_desc') rows.sort((a, b) => b.triple_digit_streak - a.triple_digit_streak);
  else if (sort === 'accel_streak_desc') rows.sort((a, b) => b.accel_streak - a.accel_streak);
  else rows.sort((a, b) => (b.latest_mktcap ?? 0) - (a.latest_mktcap ?? 0));

  document.getElementById('count').textContent = rows.length + '종목 (행 클릭하면 표)';
  document.getElementById('tbody').innerHTML = rows.slice(0, 400).map(r => `
    <tr data-code="${{r.code}}">
      <td>${{r.name}}</td><td class="sub">${{r.code}}</td><td class="sub">${{r.sector || '-'}}</td>
      <td>${{fmtMktcap(r.latest_mktcap)}}</td>
      <td class="sub">${{fmtPeriod(r.latest_period)}}${{r.latest_is_estimate ? '<span class="est-badge">추정</span>' : ''}}</td>
      <td>${{pctSpan(r.latest_yoy)}}</td>
      <td>${{pctSpan(r.prev_yoy)}}</td>
      <td>${{pctSpan(r.accel)}}</td>
      <td class="sub">${{r.triple_digit_streak > 0 ? r.triple_digit_streak + '분기' : '-'}}</td>
      <td class="sub">${{r.accel_streak > 0 ? r.accel_streak + '분기' : '-'}}</td>
    </tr>`).join('');
  document.querySelectorAll('#tbody tr').forEach(tr =>
    tr.addEventListener('click', () => openDetail(tr.dataset.code)));
}}

['fSearch','fSector','fYoyMin','fMktcapMin','fAccelOnly','fTripleOnly','fAccelStreakOnly','fSort'].forEach(id => {{
  document.getElementById(id).addEventListener('input', applyFilters);
  document.getElementById(id).addEventListener('change', applyFilters);
}});

function openDetail(code) {{
  fetch(`yoy_accel_tracker_data/${{code}}.json`).then(r => r.json()).then(d => {{
    document.getElementById('detailName').textContent = `${{d.name}} (${{d.code}})`;
    document.getElementById('detailBody').innerHTML = d.quarters.slice().reverse().map(q => `
      <tr>
        <td style="text-align:left;">${{fmtPeriod(q.period)}}${{q.is_estimate ? '<span class="est-badge">추정</span>' : ''}}</td>
        <td>${{q.op_100mil.toLocaleString()}}억</td>
        <td>${{pctSpan(q.yoy)}}</td>
      </tr>`).join('');
    document.getElementById('overlay').classList.add('open');
  }});
}}
document.getElementById('closeBtn').addEventListener('click', () => document.getElementById('overlay').classList.remove('open'));
document.getElementById('overlay').addEventListener('click', e => {{
  if (e.target.id === 'overlay') document.getElementById('overlay').classList.remove('open');
}});

applyFilters();
</script>
</body>
</html>
"""


if __name__ == "__main__":
    main()
